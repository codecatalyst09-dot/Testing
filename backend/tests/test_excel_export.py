import pytest
import openpyxl
from pathlib import Path
from fastapi.testclient import TestClient

from backend.app import app
from backend.services.a360_preprocessor import A360Preprocessor
from backend.services.a360_parser import A360Parser
from backend.services.excel_export_service import ExcelExportService

client = TestClient(app)

def test_excel_export_service_and_complexity():
    sample_a360 = {
        "name": "FinanceVendorBot",
        "variables": [
            {"name": "vFilePath", "type": "String", "input": True, "defaultValue": {"string": "C:\\Vendors.xlsx"}},
            {"name": "vRecords", "type": "Table", "input": False}
        ],
        "nodes": [
            {
                "commandName": "Excel",
                "disabled": False,
                "attributes": [
                    {"name": "action", "value": {"string": "open"}},
                    {"name": "filePath", "value": {"string": "$vFilePath$"}}
                ]
            },
            {
                "commandName": "Excel",
                "disabled": True,
                "attributes": [
                    {"name": "action", "value": {"string": "close"}}
                ]
            },
            {
                "commandName": "REST Web Service",
                "disabled": False,
                "attributes": [
                    {"name": "action", "value": {"string": "postMethod"}},
                    {"name": "url", "value": {"string": "https://vendor.api/submit"}}
                ]
            }
        ]
    }

    preprocessor = A360Preprocessor()
    cleaned, disabled = preprocessor.process(sample_a360, "FinanceVendorBot", "FinanceVendorBot.json")

    parser = A360Parser()
    wf = parser.parse(
        cleaned_taskbots_data=[("FinanceVendorBot", "FinanceVendorBot.json", cleaned)],
        raw_data_map={"FinanceVendorBot": sample_a360},
        disabled_actions=disabled,
        workflow_id="wf-test-excel",
        workflow_name="Finance Vendor Bot"
    )

    # Validate stats and user's complexity rule (< 200 = Easy)
    assert wf.statistics.totalActions == 2
    assert wf.statistics.totalDisabledActions == 1
    assert wf.statistics.totalStepsEvaluated == 3
    assert wf.statistics.migrationComplexity == "Easy"

    test_excel_path = Path("storage/test_finance_migration.xlsx")
    ExcelExportService.generate_migration_excel(wf, test_excel_path)
    assert test_excel_path.exists()

    # Inspect Workbook
    wb = openpyxl.load_workbook(test_excel_path, data_only=True)
    assert "Migration Action Mapping" in wb.sheetnames
    assert "Variables Dictionary" in wb.sheetnames
    assert "Disabled Actions Audit" in wb.sheetnames
    assert "Package Reference Summary" in wb.sheetnames

    ws_steps = wb["Migration Action Mapping"]
    # Check top complexity banner
    assert "EASY MIGRATION" in str(ws_steps["A5"].value)

    # Check that both active and disabled steps are present
    step_rows = []
    for r in range(9, ws_steps.max_row + 1):
        step_rows.append({
            "step": ws_steps.cell(r, 1).value,
            "status": ws_steps.cell(r, 2).value,
            "package": ws_steps.cell(r, 4).value,
            "rec_action": ws_steps.cell(r, 11).value
        })

    assert len(step_rows) == 3
    active_rows = [s for s in step_rows if s["status"] == "ACTIVE"]
    disabled_rows = [s for s in step_rows if s["status"] == "DISABLED"]
    assert len(active_rows) == 2
    assert len(disabled_rows) == 1
    assert disabled_rows[0]["package"] == "Excel Advanced"

def test_download_excel_api_endpoint(tmp_path):
    # Upload sample bot through API
    payload = {
        "name": "ApiUploadBot",
        "nodes": [
            {
                "commandName": "Excel",
                "disabled": False,
                "attributes": [
                    {"name": "action", "value": {"string": "open"}}
                ]
            },
            {
                "commandName": "Excel",
                "disabled": True,
                "attributes": [
                    {"name": "action", "value": {"string": "save"}}
                ]
            }
        ]
    }

    import json
    from backend.services.pipeline_orchestrator import PipelineOrchestrator
    json_bytes = json.dumps(payload).encode("utf-8")
    resp = client.post(
        "/api/upload",
        files={"file": ("ApiUploadBot.json", json_bytes, "application/json")}
    )
    assert resp.status_code == 200
    job_id = resp.json()["job_id"]

    # Trigger analysis synchronously
    PipelineOrchestrator.run_analysis(job_id)

    # Test GET /api/jobs/{job_id}/download-excel
    dl_resp = client.get(f"/api/jobs/{job_id}/download-excel")
    assert dl_resp.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in dl_resp.headers["content-type"]
    assert len(dl_resp.content) > 1000  # Valid binary Excel file
