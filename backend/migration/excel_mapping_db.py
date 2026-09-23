import os
import openpyxl
from pathlib import Path
from typing import Dict, Any, Optional, Tuple, List

from backend.utils.paths import get_mapping_file_path

MAPPING_FILE_PATH = get_mapping_file_path()

class ExcelMappingDB:
    _instance = None
    _mappings: List[Dict[str, Any]] = []
    _lookup: Dict[str, Dict[str, Any]] = {}

    @property
    def mappings(self) -> List[Dict[str, Any]]:
        return self._mappings

    @classmethod
    def get_instance(cls):
        if cls._instance is None:
            cls._instance = cls()
            cls._instance._load_mappings()
        return cls._instance

    def _load_mappings(self):
        self._mappings = []
        self._lookup = {}

        if not MAPPING_FILE_PATH.exists():
            return

        try:
            wb = openpyxl.load_workbook(MAPPING_FILE_PATH, data_only=True)
            if "Mapping" not in wb.sheetnames:
                return
            ws = wb["Mapping"]
            headers = [cell.value for cell in ws[1]]
            
            # Map column indices
            pkg_col = headers.index("AA Package") if "AA Package" in headers else 0
            act_col = headers.index("AA Action") if "AA Action" in headers else 1
            desc_col = headers.index("AA Action Description") if "AA Action Description" in headers else 2
            
            # Find PAD Category, PAD Equivalent, PA Cloud Equivalent, Notes
            pad_cat_col = next((i for i, h in enumerate(headers) if h and "PAD Category" in str(h) or "Desktop" in str(h) and "Category" in str(h)), 3)
            pad_eq_col = next((i for i, h in enumerate(headers) if h and "PAD Equivalent" in str(h) or "Desktop" in str(h) and "Equivalent" in str(h)), 4)
            cloud_eq_col = next((i for i, h in enumerate(headers) if h and "Cloud" in str(h) and "Equivalent" in str(h)), 5)
            notes_col = next((i for i, h in enumerate(headers) if h and "Migration Notes" in str(h) or "Notes" in str(h)), 6)

            for r in range(2, ws.max_row + 1):
                vals = [cell.value for cell in ws[r]]
                if not any(v is not None for v in vals):
                    continue

                pkg = str(vals[pkg_col]).strip() if vals[pkg_col] is not None else ""
                act = str(vals[act_col]).strip() if vals[act_col] is not None else ""
                desc = str(vals[desc_col]).strip() if desc_col < len(vals) and vals[desc_col] is not None else ""
                pad_cat = str(vals[pad_cat_col]).strip() if pad_cat_col < len(vals) and vals[pad_cat_col] is not None else ""
                pad_act = str(vals[pad_eq_col]).strip() if pad_eq_col < len(vals) and vals[pad_eq_col] is not None else ""
                cloud_act = str(vals[cloud_eq_col]).strip() if cloud_eq_col < len(vals) and vals[cloud_eq_col] is not None else ""
                notes = str(vals[notes_col]).strip() if notes_col < len(vals) and vals[notes_col] is not None else ""

                entry = {
                    "aa_package": pkg,
                    "aa_action": act,
                    "aa_description": desc,
                    "pad_category": pad_cat,
                    "pad_action": pad_act,
                    "cloud_action": cloud_act,
                    "migration_notes": notes,
                }
                self._mappings.append(entry)

                # Store key variations for robust lookup
                key1 = f"{pkg.lower()}|{act.lower()}"
                self._lookup[key1] = entry

                # Package only fallback
                pkg_key = pkg.lower()
                if pkg_key not in self._lookup:
                    self._lookup[pkg_key] = entry

            # Add supplementary catalog for packages not in raw Excel export
            self._load_supplementary_catalog()

        except Exception as e:
            print(f"Error loading Excel mapping: {e}")

    def _load_supplementary_catalog(self):
        supplementary = [
            # SAP GUI
            {"aa_package": "SAP", "aa_action": "connect", "aa_description": "Connect to SAP GUI server session", "pad_category": "SAP GUI Automation", "pad_action": "SAP GUI Automation / Connect to SAP", "cloud_action": "No direct equivalent", "migration_notes": "Requires SAP GUI client and scripting enabled on server."},
            {"aa_package": "SAP", "aa_action": "runTransaction", "aa_description": "Execute transaction code in SAP GUI", "pad_category": "SAP GUI Automation", "pad_action": "SAP GUI Automation (Execute Transaction)", "cloud_action": "No direct equivalent", "migration_notes": "Maps to PAD SAP execute transaction code action."},
            {"aa_package": "SAP", "aa_action": "setText", "aa_description": "Set text in SAP GUI element", "pad_category": "SAP GUI Automation", "pad_action": "Populate text field in SAP window", "cloud_action": "No direct equivalent", "migration_notes": "Maps to PAD SAP GUI field population."},
            {"aa_package": "SAP", "aa_action": "press", "aa_description": "Press button in SAP GUI window", "pad_category": "SAP GUI Automation", "pad_action": "Press button in SAP window", "cloud_action": "No direct equivalent", "migration_notes": "Maps to PAD SAP GUI button press."},
            {"aa_package": "SAP", "aa_action": "getText", "aa_description": "Get text from SAP GUI element", "pad_category": "SAP GUI Automation", "pad_action": "Get details of element on SAP window", "cloud_action": "No direct equivalent", "migration_notes": "Extracts text from SAP window into flow variable."},
            {"aa_package": "SAP", "aa_action": "selectItem", "aa_description": "Select item in SAP menu or table", "pad_category": "SAP GUI Automation", "pad_action": "Select menu item in SAP window", "cloud_action": "No direct equivalent", "migration_notes": "Selects item in SAP table or toolbar."},
            {"aa_package": "SAP", "aa_action": "closeSession", "aa_description": "Close active SAP session", "pad_category": "SAP GUI Automation", "pad_action": "Close SAP session", "cloud_action": "No direct equivalent", "migration_notes": "Gracefully terminates SAP GUI session."},

            # Word
            {"aa_package": "Word", "aa_action": "open", "aa_description": "Open Word document", "pad_category": "Word", "pad_action": "Launch Word", "cloud_action": "Word Online (Business) Connector", "migration_notes": "Desktop Word manipulation maps to PAD Word actions."},
            {"aa_package": "Word", "aa_action": "save", "aa_description": "Save Word document", "pad_category": "Word", "pad_action": "Save Word document", "cloud_action": "Word Online (Business) Connector", "migration_notes": "Saves active Word document."},
            {"aa_package": "Word", "aa_action": "close", "aa_description": "Close Word document", "pad_category": "Word", "pad_action": "Close Word", "cloud_action": "Word Online (Business) Connector", "migration_notes": "Closes Word document."},
            {"aa_package": "Word", "aa_action": "read", "aa_description": "Read text from Word document", "pad_category": "Word", "pad_action": "Read from Word document", "cloud_action": "Word Online (Business) Connector", "migration_notes": "Extracts text from Word document."},
            {"aa_package": "Word", "aa_action": "replaceText", "aa_description": "Replace text in Word document", "pad_category": "Word", "pad_action": "Find and replace text in Word document", "cloud_action": "Word Online (Business) Connector", "migration_notes": "Replaces text strings in document."},

            # SharePoint Online
            {"aa_package": "SharePoint", "aa_action": "getFile", "aa_description": "Download or read SharePoint file", "pad_category": "Cloud Storage", "pad_action": "Download file from SharePoint", "cloud_action": "SharePoint: Get file content", "migration_notes": "Cloud-native SharePoint connector."},
            {"aa_package": "SharePoint", "aa_action": "createFile", "aa_description": "Upload file to SharePoint", "pad_category": "Cloud Storage", "pad_action": "Upload file to SharePoint", "cloud_action": "SharePoint: Create file", "migration_notes": "Cloud-native SharePoint connector."},
            {"aa_package": "SharePoint", "aa_action": "getItems", "aa_description": "Get items from SharePoint list", "pad_category": "Cloud Storage", "pad_action": "Get SharePoint list items", "cloud_action": "SharePoint: Get items", "migration_notes": "Cloud-native SharePoint connector."},
            {"aa_package": "SharePoint", "aa_action": "updateProperties", "aa_description": "Update SharePoint file metadata", "pad_category": "Cloud Storage", "pad_action": "Update SharePoint properties", "cloud_action": "SharePoint: Update file properties", "migration_notes": "Cloud-native SharePoint connector."},

            # OneDrive for Business
            {"aa_package": "OneDrive", "aa_action": "getFile", "aa_description": "Get file content from OneDrive", "pad_category": "Cloud Storage", "pad_action": "Download file from OneDrive", "cloud_action": "OneDrive for Business: Get file content", "migration_notes": "Cloud-native OneDrive connector."},
            {"aa_package": "OneDrive", "aa_action": "createFile", "aa_description": "Upload or create file in OneDrive", "pad_category": "Cloud Storage", "pad_action": "Upload file to OneDrive", "cloud_action": "OneDrive for Business: Create file", "migration_notes": "Cloud-native OneDrive connector."},

            # Microsoft Teams
            {"aa_package": "Teams", "aa_action": "postMessage", "aa_description": "Post message to channel or chat", "pad_category": "Collaboration", "pad_action": "Post message in Teams", "cloud_action": "Microsoft Teams: Post message in a chat or channel", "migration_notes": "Cloud-native Teams connector."},
            {"aa_package": "Teams", "aa_action": "postAdaptiveCard", "aa_description": "Post interactive adaptive card", "pad_category": "Collaboration", "pad_action": "Post adaptive card in Teams", "cloud_action": "Microsoft Teams: Post adaptive card and wait for a response", "migration_notes": "Cloud-native Teams interactive card."},

            # Zip / Archive
            {"aa_package": "Zip", "aa_action": "compress", "aa_description": "Compress files into ZIP archive", "pad_category": "File", "pad_action": "ZIP files", "cloud_action": "No native action  use OneDrive Extract or Encodian", "migration_notes": "Local file compression maps to PAD ZIP files action."},
            {"aa_package": "Zip", "aa_action": "uncompress", "aa_description": "Extract files from ZIP archive", "pad_category": "File", "pad_action": "Unzip files", "cloud_action": "OneDrive for Business: Extract folder", "migration_notes": "Local file extraction maps to PAD Unzip files action."},

            # Cryptography
            {"aa_package": "Cryptography", "aa_action": "encrypt", "aa_description": "Encrypt text with AES or RSA", "pad_category": "Cryptography", "pad_action": "Encrypt text with AES", "cloud_action": "Azure Key Vault / Azure Function", "migration_notes": "Desktop AES encryption action in PAD."},
            {"aa_package": "Cryptography", "aa_action": "decrypt", "aa_description": "Decrypt text with AES or RSA", "pad_category": "Cryptography", "pad_action": "Decrypt text with AES", "cloud_action": "Azure Key Vault / Azure Function", "migration_notes": "Desktop AES decryption action in PAD."},
            # Log to File / Logging
            {"aa_package": "Log to file", "aa_action": "log", "aa_description": "Append text entry to log file", "pad_category": "File", "pad_action": "Write text to file", "cloud_action": "OneDrive / SharePoint: Append to file (or Dataverse log / Azure Application Insights)", "migration_notes": "Cloud-native logging maps to SharePoint/OneDrive append or Dataverse log row."},
            {"aa_package": "Log to file", "aa_action": "write", "aa_description": "Write text entry to log file", "pad_category": "File", "pad_action": "Write text to file", "cloud_action": "OneDrive / SharePoint: Append to file (or Dataverse log / Azure Application Insights)", "migration_notes": "Cloud-native logging maps to SharePoint/OneDrive append or Dataverse log row."},
            {"aa_package": "Log", "aa_action": "log", "aa_description": "Append text entry to log file", "pad_category": "File", "pad_action": "Write text to file", "cloud_action": "OneDrive / SharePoint: Append to file (or Dataverse log / Azure Application Insights)", "migration_notes": "Cloud-native logging maps to SharePoint/OneDrive append or Dataverse log row."},
            {"aa_package": "Log to file", "aa_action": "", "aa_description": "Append text entry to log file", "pad_category": "File", "pad_action": "Write text to file", "cloud_action": "OneDrive / SharePoint: Append to file (or Dataverse log / Azure Application Insights)", "migration_notes": "Cloud-native logging maps to SharePoint/OneDrive append or Dataverse log row."},
        ]

        for item in supplementary:
            self._mappings.append(item)
            key = f"{item['aa_package'].lower()}|{item['aa_action'].lower()}"
            self._lookup[key] = item
            pkg_key = item['aa_package'].lower()
            if pkg_key not in self._lookup:
                self._lookup[pkg_key] = item

    def find_mapping(self, command: str, operation: Optional[str] = None) -> Dict[str, Any]:
        """
        Lookup mapping from the official Excel database + supplementary catalog.
        """
        cmd_clean = (command or "").strip()
        op_clean = (operation or "").strip()
        cmd_lower = cmd_clean.lower()
        op_lower = op_clean.lower()

        # Command aliases mapping to official AA Packages
        alias_map = {
            "excel": "Excel Advanced",
            "excelbasic": "Excel Basic",
            "exceladvanced": "Excel Advanced",
            "msexcel": "Excel Advanced",
            "word": "Word",
            "msword": "Word",
            "runtask": "Task Bot",
            "subtask": "Task Bot",
            "taskbot": "Task Bot",
            "http": "REST Web Service",
            "rest": "REST Web Service",
            "restapi": "REST Web Service",
            "restwebservice": "REST Web Service",
            "webapi": "REST Web Service",
            "soap": "SOAP Web Service",
            "soapwebservice": "SOAP Web Service",
            "sendemail": "Email",
            "mail": "Email",
            "outlook": "Email",
            "exchange": "Email",
            "csv": "CSV/TXT",
            "txt": "CSV/TXT",
            "textfile": "CSV/TXT",
            "datatable": "Data Table",
            "table": "Data Table",
            "webautomation": "Browser",
            "capture": "Recorder",
            "recorder": "Recorder",
            "objectcloning": "Recorder",
            "messagebox": "Message box",
            "msgbox": "Message box",
            "prompt": "Prompt",
            "python": "Python Script",
            "pythonscript": "Python Script",
            "vbs": "VBScript",
            "vbscript": "VBScript",
            "powershell": "PowerShell",
            "posh": "PowerShell",
            "javascript": "JavaScript",
            "dll": "DLL",
            "string": "String",
            "number": "Number",
            "boolean": "Boolean",
            "datetime": "DateTime",
            "date": "DateTime",
            "loop": "Loop",
            "while": "Loop",
            "if": "If",
            "condition": "If",
            "delay": "Delay",
            "wait": "Delay",
            "file": "File",
            "folder": "Folder",
            "filesystem": "File",
            "database": "Database",
            "sql": "Database",
            "activedirectory": "Active Directory",
            "clipboard": "Clipboard",
            "ocr": "OCR",
            "pdf": "PDF",
            "window": "Window",
            "windows": "Window",
            "application": "Application",
            "process": "Application",
            "system": "System",
            "doscommand": "System",
            "mouse": "Mouse",
            "keystroke": "Simulate Keystrokes",
            "keystrokes": "Simulate Keystrokes",
            "simulatekeystrokes": "Simulate Keystrokes",
            "sendkeys": "Simulate Keystrokes",
            "xml": "XML",
            "json": "JSON",
            "sap": "SAP",
            "sapgui": "SAP",
            "workload": "Workload",
            "queue": "Workload",
            "iqbot": "OCR",
            "iq bot": "OCR",
            "errorhandler": "Error Handler",
            "error handling": "Error Handler",
            "trycatch": "Error Handler",
            "sharepoint": "SharePoint",
            "onedrive": "OneDrive",
            "teams": "Teams",
            "msteams": "Teams",
            "googledrive": "Google Drive",
            "googlesheets": "Google Sheets",
            "zip": "Zip",
            "archive": "Zip",
            "cryptography": "Cryptography",
            "crypto": "Cryptography",
            "service": "Service",
            "printer": "Printer",
            "log": "Log to file",
            "logtofile": "Log to file",
            "logfile": "Log to file",
        }

        matched_pkg = alias_map.get(cmd_lower.replace(" ", "").replace("_", ""), cmd_clean)

        # 1. Exact match package|action
        key = f"{matched_pkg.lower()}|{op_lower}"
        if key in self._lookup:
            return self._format_result(self._lookup[key], cmd_clean, op_clean)

        # 2. Match by package with similar action
        pkg_lower = matched_pkg.lower()
        for item in self._mappings:
            if item["aa_package"].lower() == pkg_lower:
                if op_lower and op_lower in item["aa_action"].lower():
                    return self._format_result(item, cmd_clean, op_clean)

        # 3. Match by package general row
        for item in self._mappings:
            if item["aa_package"].lower() == pkg_lower:
                return self._format_result(item, cmd_clean, op_clean)

        # 4. Keyword search across packages in database
        for item in self._mappings:
            if cmd_lower and cmd_lower in item["aa_package"].lower():
                return self._format_result(item, cmd_clean, op_clean)

        # 5. Unknown command policy -> Manual Review
        return {
            "aa_package": cmd_clean,
            "aa_action": op_clean or "Execute",
            "aa_description": "Custom or proprietary Automation Anywhere action.",
            "pad_category": "Manual Review Required",
            "pad_action": f"Manual Review ({cmd_clean})",
            "cloud_action": "No direct equivalent",
            "platform": "Manual Review",
            "recommended_action": f"Manual Review: {cmd_clean}",
            "migration_notes": "Action not found in standard A360 mapping reference. Review original bot logic and select equivalent connector or script."
        }

    def _format_result(self, raw_entry: Dict[str, Any], orig_cmd: str, orig_op: str) -> Dict[str, Any]:
        pad_action = raw_entry.get("pad_action", "")
        cloud_action = raw_entry.get("cloud_action", "")
        notes = raw_entry.get("migration_notes") or ""

        # Determine target platform based on Cloud-First principle:
        # Check if cloud action is unavailable or unsupported
        is_cloud_unsupported = (
            not cloud_action or
            "no direct" in cloud_action.lower() or
            "no native" in cloud_action.lower() or
            cloud_action.strip().lower() in ("none", "n/a", "not supported")
        )

        # Packages that strictly require on-premises desktop client or local OS
        desktop_only_packages = [
            "sap", "sapgui", "application", "browser", "recorder", "window",
            "mouse", "simulate keystrokes", "clipboard", "terminal emulator",
            "citrix", "python script", "vbscript", "powershell", "system",
            "printer", "dll", "excel advanced", "excel basic", "word"
        ]
        pkg_lower = raw_entry["aa_package"].lower()
        is_web_service = "web service" in pkg_lower or "rest" in pkg_lower or "soap" in pkg_lower or "http" in pkg_lower
        is_desktop_only = False
        if not is_web_service:
            if any(dp in pkg_lower for dp in desktop_only_packages) or pkg_lower == "service":
                is_desktop_only = True

        # Hybrid orchestration
        if orig_cmd.lower() in ("runtask", "taskbot", "subtask") or raw_entry["aa_package"] == "Task Bot":
            platform = "Hybrid"
            rec_action = "Run a flow built with Power Automate for desktop (or Run Child Flow)"
        # Desktop-bound packages
        elif is_desktop_only:
            platform = "Power Automate Desktop"
            rec_action = pad_action or f"PAD Desktop Action ({raw_entry['aa_package']})"
        # Cloud First: if cloud equivalent is supported, select Cloud
        elif not is_cloud_unsupported:
            platform = "Power Automate Cloud"
            rec_action = cloud_action
        # Fallback to Desktop if Cloud is not supported
        elif pad_action and "no direct" not in pad_action.lower() and pad_action.lower() not in ("none", "n/a"):
            platform = "Power Automate Desktop"
            rec_action = pad_action
        # Fallback to Manual Review if neither Cloud nor Desktop is supported
        else:
            platform = "Manual Review"
            rec_action = f"Manual Review: {raw_entry['aa_package']}"

        return {
            "aa_package": raw_entry.get("aa_package", orig_cmd),
            "aa_action": orig_op or raw_entry.get("aa_action", ""),
            "aa_description": raw_entry.get("aa_description", ""),
            "pad_category": raw_entry.get("pad_category", "Custom"),
            "pad_action": pad_action,
            "cloud_action": cloud_action,
            "platform": platform,
            "recommended_action": rec_action,
            "migration_notes": notes
        }
